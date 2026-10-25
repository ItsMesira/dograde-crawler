"""ส่งออกผลการเรียนเป็นไฟล์ Excel (เลือกได้ว่าปุ่มไหนใน 6 ปุ่ม)

อ่านรายชื่อจาก CSV/XLSX แล้วล็อกอินทีละคนด้วยเอนจินเดิมของ grade_watch_v3
ไม่แตะ state.json / baseline.png จึงไม่กระทบการแจ้งเตือนของบอท
และไม่ยิง Discord เลย (ตั้ง g.WEBHOOK = "")

ติดตั้งเพิ่มครั้งเดียว:
    pip install openpyxl

ตัวอย่าง:
    python export_grades.py                              # เมนูเลือกภาค -> exports/
    python export_grades.py --term 1                     # ปี1ภาค1 ตรง ๆ
    python export_grades.py --term 0                     # ทั้ง 6 ภาค
    python export_grades.py --dry-run                    # ตรวจรายชื่อก่อน (ไม่เปิดเว็บ)
    python export_grades.py --limit 1 --headed           # ทดลองคนแรก แบบเห็นหน้าจอ
    python export_grades.py --terms 1-1,1-2              # หลายภาค (คั่นด้วย comma)
    python export_grades.py --roster รายชื่อ.xlsx        # ใช้ไฟล์ Excel แทน CSV
"""
import argparse
import asyncio
import csv
import datetime
import re
import sys
import time
from pathlib import Path

ROOT = Path(__file__).resolve().parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

try:  # กัน cmd ภาษาไทยพังตอนพิมพ์อีโมจิ
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

import grade_watch_v3 as g

g.WEBHOOK = ""                  # ห้ามยิง Discord ตอนส่งออก
g.BASE_DIR = ROOT / "shots"     # ให้อยู่ที่ repo root เสมอ ไม่ขึ้นกับ CWD

DEFAULT_ROSTER = ROOT / "students.csv"
EXPORT_DIR = ROOT / "exports"
# ปุ่มบนเว็บเรียงตามนี้: 1=ปี1ภาค1 ... 6=ปี3ภาค2
TERM_BUTTONS = [(1, 1), (1, 2), (2, 1), (2, 2), (3, 1), (3, 2)]
DEFAULT_TERM = (1, 1)

# ชื่อคอลัมน์ที่ยอมรับ (ไทย/อังกฤษ) — ใช้ตัวพิมพ์เล็กและตัดช่องว่างก่อนเทียบ
ALIASES = {
    "sid": ("sid", "student_id", "studentid", "id", "รหัส", "รหัสนักเรียน",
            "เลขประจำตัว", "เลขประจำตัวนักเรียน"),
    "birthday": ("birthday", "birth", "dob", "pw", "password", "วันเกิด",
                 "วันเดือนปีเกิด"),
    "name": ("name", "ชื่อ", "ชื่อ-สกุล", "ชื่อสกุล", "ชื่อนักเรียน"),
    "room": ("room", "ห้อง", "ห้องเรียน"),
    "seat": ("seat", "no", "เลขที่", "เลขที่นั่ง"),
}


# ---------------- รายชื่อ ----------------
def _key(s):
    return re.sub(r"\s+", "", str(s or "")).lower()


def _map_header(h):
    k = _key(h)
    for field, names in ALIASES.items():
        if k in {_key(n) for n in names}:
            return field
    return None


def _read_csv(path):
    for enc in ("utf-8-sig", "cp874", "tis-620"):
        try:
            with path.open(encoding=enc, newline="") as f:
                return [row for row in csv.reader(f)]
        except UnicodeDecodeError:
            continue
    raise SystemExit(f"อ่าน {path.name} ไม่ได้ — บันทึกเป็น UTF-8 แล้วลองใหม่")


def _read_xlsx(path):
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise SystemExit("ยังไม่ได้ติดตั้ง openpyxl — รัน: pip install openpyxl")
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = [list(row) for row in ws.iter_rows(values_only=True)]
    wb.close()
    return rows


def _norm_birthday(v):
    """คืน วว/ดด/ปปปป (พ.ศ.) หรือ None ถ้ารูปแบบไม่ถูก"""
    if isinstance(v, datetime.datetime):
        v = v.date()
    if isinstance(v, datetime.date):
        return f"{v.day:02d}/{v.month:02d}/{v.year}"
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        try:  # Excel serial date
            d = datetime.datetime(1899, 12, 30) + datetime.timedelta(days=float(v))
            return f"{d.day:02d}/{d.month:02d}/{d.year}"
        except Exception:
            return None
    s = str(v or "").strip()
    m = re.fullmatch(r"(\d{4})-(\d{1,2})-(\d{1,2})", s)
    if m:
        y, mo, d = map(int, m.groups())
        return f"{d:02d}/{mo:02d}/{y}" if 2400 <= y <= 2699 else None
    m = re.fullmatch(r"(\d{1,2})\s*[/.\-]\s*(\d{1,2})\s*[/.\-]\s*(\d{4})", s)
    if not m:
        return None
    d, mo, y = map(int, m.groups())
    if not (1 <= d <= 31 and 1 <= mo <= 12 and 2400 <= y <= 2699):
        return None
    return f"{d:02d}/{mo:02d}/{y}"


def _text(v):
    s = str(v if v is not None else "").strip()
    if s.endswith(".0") and s[:-2].isdigit():
        s = s[:-2]
    return s


def load_roster(path):
    path = Path(path)
    if not path.exists():
        raise SystemExit(
            f"ไม่พบไฟล์รายชื่อ {path}\n"
            f"คัดลอก students.example.csv เป็น students.csv แล้วใส่ข้อมูลจริง")
    rows = _read_xlsx(path) if path.suffix.lower() in (".xlsx", ".xlsm") else _read_csv(path)
    rows = [r for r in rows if any(_text(c) for c in r)]
    if not rows:
        raise SystemExit(f"{path.name} ว่างเปล่า")

    cols = {}
    for i, h in enumerate(rows[0]):
        field = _map_header(h)
        if field and field not in cols:
            cols[field] = i
    missing = [f for f in ("sid", "birthday") if f not in cols]
    if missing:
        raise SystemExit(
            f"ไม่พบคอลัมน์ {missing} ใน {path.name}\n"
            f"หัวตารางที่อ่านได้: {[str(h) for h in rows[0]]}\n"
            f"ต้องมีอย่างน้อย: sid,birthday")

    out, seen, errors = [], set(), []
    for n, row in enumerate(rows[1:], start=2):
        def cell(field):
            i = cols.get(field)
            return row[i] if i is not None and i < len(row) else ""

        sid = _text(cell("sid"))
        raw_bd = cell("birthday")
        bd = _norm_birthday(raw_bd)
        if not sid:
            errors.append(f"แถว {n}: ไม่มีรหัสนักเรียน")
            continue
        if sid in seen:
            errors.append(f"แถว {n}: รหัส {sid} ซ้ำ")
            continue
        if not bd:
            errors.append(
                f"แถว {n}: วันเกิด '{_text(raw_bd)}' ต้องเป็น วว/ดด/ปปปป พ.ศ. "
                f"เช่น 01/01/2550")
            continue
        seen.add(sid)
        out.append({"sid": sid, "pw": bd,
                    "name": _text(cell("name")),
                    "room": _text(cell("room")),
                    "seat": _text(cell("seat"))})
    return out, errors


def parse_terms(s):
    if not s:
        return [DEFAULT_TERM]
    out = []
    for part in s.split(","):
        m = re.fullmatch(r"\s*(\d+)\s*[-/]\s*(\d+)\s*", part)
        if not m:
            raise SystemExit(f"รูปแบบ --terms ไม่ถูก: {part} (ใช้ เช่น 1-1 หรือ 1-1,1-2)")
        y, t = map(int, m.groups())
        if (y, t) not in out:
            out.append((y, t))
    return sorted(out)


def parse_term_arg(s):
    """เลขปุ่ม 1-6 (0 = ทั้ง 6 ภาค) -> list ของ (ปี, ภาค)"""
    try:
        n = int(str(s).strip())
    except ValueError:
        raise SystemExit(f"--term ต้องเป็นเลข 0-6 (ได้รับ {s!r})")
    if n == 0:
        return list(TERM_BUTTONS)
    if 1 <= n <= 6:
        return [TERM_BUTTONS[n - 1]]
    raise SystemExit(f"--term ต้องเป็นเลข 0-6 (ได้รับ {n})")


def choose_term():
    """เมนูเลือกปี/ภาค — ถ้าไม่ใช่ terminal (เช่น pipe/CI) ใช้ค่าเริ่มต้น 1-1"""
    if not sys.stdin.isatty():
        print("[i] ไม่ได้รันใน terminal — ใช้ค่าเริ่มต้น ปี1ภาค1 (ระบุ --term เพื่อเลือก)")
        return [DEFAULT_TERM]

    print("\nเลือกภาคที่จะดึง:")
    for i, (y, t) in enumerate(TERM_BUTTONS, 1):
        print(f"  {i}) ปี{y} ภาค{t}")
    print("  0) ทั้ง 6 ภาค")
    while True:
        try:
            raw = input("เลือก [0-6] (Enter = 1): ").strip()
        except (EOFError, KeyboardInterrupt):
            print()
            return [DEFAULT_TERM]
        if not raw:
            return [DEFAULT_TERM]
        try:
            n = int(raw)
        except ValueError:
            print("  พิมพ์เป็นตัวเลข 0-6")
            continue
        if n == 0:
            return list(TERM_BUTTONS)
        if 1 <= n <= 6:
            return [TERM_BUTTONS[n - 1]]
        print("  พิมพ์เป็นตัวเลข 0-6")


def print_roster(roster, errors, terms, total):
    print(f"{'#':>3}  {'รหัส':<10} {'วันเกิด':<12} {'ห้อง':<5} {'เลขที่':<6} ชื่อ")
    for i, r in enumerate(roster, 1):
        print(f"{i:>3}  {r['sid']:<10} {r['pw']:<12} {r['room']:<5} {r['seat']:<6} {r['name']}")
    print(f"\nใช้ได้ {len(roster)} คน" + (f" (แสดง {len(roster)}/{total})" if len(roster) < total else ""))
    if errors:
        print(f"ใช้ไม่ได้ {len(errors)} แถว:")
        for e in errors:
            print("  -", e)
    print("จะดึง: " + ", ".join(f"ปี{y}ภาค{t}" for y, t in terms))


# ---------------- ดึงข้อมูล ----------------
async def _scrape_student(browser, rec, terms, args):
    sid = rec["sid"]
    res = {"sid": sid, "roster_name": rec.get("name", ""), "student": {},
           "terms": {}, "error": None}
    S = {"sid": sid, "pw": rec["pw"], "name": rec.get("name") or sid,
         "state": {"terms": {}}, "watch_terms": terms, "watch_pp1": False}
    sess = g.StudentSession(browser, S)
    try:
        await sess.open()
        await sess.login()
        for (y, t) in terms:
            key = f"{y}/{t}"
            try:
                method = await g.go_term(sess.page, y, t)
                if method is None:
                    # ปุ่มไม่มีบนหน้า = ภาคนี้ยังไม่เปิด ให้ข้ามแบบไม่นับเป็น error
                    res["terms"][key] = {"error": "ไม่พบปุ่มปี/ภาคบนหน้าเว็บ",
                                         "skipped": True}
                    continue
                data = await g.read_term_table(sess.page)

                # กันข้อมูลภาคอื่นถูกเขียนทับ: ปุ่มปี1..3 = ม.4..6
                level = data.get("level")
                if level is not None and level != 3 + y:
                    res["terms"][key] = {
                        "error": f"หน้าเป็น ม.{level} ไม่ใช่ ปี{y} (คาดว่า ม.{3 + y})",
                        "skipped": True}
                    continue
                if data.get("term") is not None and data["term"] != t:
                    res["terms"][key] = {
                        "error": f"หน้าเป็นภาค {data['term']} ไม่ใช่ภาค {t}",
                        "skipped": True}
                    continue

                stu = data.get("student") or {}
                stu["sid"] = stu.get("sid") or sid
                stu["name"] = stu.get("name") or S["name"]
                if not res["student"] and (stu.get("name") or stu.get("room")):
                    res["student"] = stu
                data["student"] = stu
                res["terms"][key] = data

                if args.save_html:
                    d = EXPORT_DIR / "html"
                    d.mkdir(parents=True, exist_ok=True)
                    (d / f"{sid}_Y{y}T{t}.html").write_text(
                        await sess.page.content(), encoding="utf-8")
                if args.shots:
                    folder = g.special_folder(sid, f"export_Y{y}T{t}")
                    await g.snap(sess.page, folder, f"Y{y}T{t}")
            except Exception as e:
                res["terms"][key] = {"error": str(e)[:160]}
        return res
    except Exception as e:
        res["error"] = str(e)[:200]
        return res
    finally:
        try:
            await sess.close()
        except Exception:
            pass


def _print_result(r, done, total):
    status, note = _status(r)
    mark = "✔" if status == "ok" else "✘"
    print(f"[{done:>2}/{total}] {r['sid']} {mark} "
          f"{('— ' + note) if note else ''}".rstrip())


def _transient(res):
    """ข้อผิดพลาดชั่วคราว (เว็บ 5xx/หมดเวลา) — ลองใหม่ได้"""
    msg = (res.get("error") or "").lower()
    return any(k in msg for k in ("http 5", "timeout", "timed out", "net::",
                                  "err_", "econn", "target closed"))


async def _run(roster, terms, args):
    results = []
    by_roster = {r["sid"]: r for r in roster}
    async with g.async_playwright() as p:
        browser = await p.chromium.launch(headless=not args.headed)
        try:
            for i in range(0, len(roster), args.workers):
                chunk = roster[i:i + args.workers]
                if i:
                    await asyncio.sleep(args.stagger)
                done = await asyncio.gather(
                    *(_scrape_student(browser, r, terms, args) for r in chunk))
                for r in done:
                    results.append(r)
                    _print_result(r, len(results), len(roster))

            # ลองใหม่เฉพาะที่ล้มเหลวเพราะเว็บ 5xx/หมดเวลา (ไม่ลองซ้ำถ้ารหัส/วันเกิดผิด)
            for attempt in range(1, args.retries + 1):
                failed = [r for r in results if _transient(r)]
                if not failed:
                    break
                print(f"\n↻ ลองใหม่รอบ {attempt}: {len(failed)} คน (ทีละคน)")
                index = {r["sid"]: i for i, r in enumerate(results)}
                for r in failed:
                    await asyncio.sleep(args.stagger)
                    new = await _scrape_student(browser, by_roster[r["sid"]], terms, args)
                    results[index[r["sid"]]] = new
                    _print_result(new, index[r["sid"]] + 1, len(roster))
        finally:
            await browser.close()
    return results


def _status(res):
    if not res:
        return "ไม่ได้รัน", ""
    if res.get("error"):
        return "ล้มเหลว", res["error"]
    errs = [f"{k}: {v['error']}" for k, v in res["terms"].items()
            if "error" in v and not v.get("skipped")]
    if errs:
        return "ล้มเหลว", "; ".join(errs)[:300]
    if not any((v.get("subjects") or []) for v in res["terms"].values()):
        return "ไม่มีข้อมูล", "ล็อกอินได้แต่ไม่พบตารางเกรด"
    return "ok", ""


# ---------------- เขียน Excel ----------------
def _safe_sheet(name):
    return re.sub(r"[\[\]:*?/\\]", "-", name)[:31]


def _style_header(ws, row, ncols):
    from openpyxl.styles import Alignment, Font, PatternFill
    fill = PatternFill("solid", fgColor="3F4E5E")
    font = Font(bold=True, color="FFFFFF", name="Tahoma", size=10)
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 28


def _set_widths(ws, widths):
    from openpyxl.utils import get_column_letter
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = min(max(w, 8), 45)


def _label(level):
    return f"ม.{level}" if level else ""


def write_workbook(results, roster, terms, out_path):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font
    except ImportError:
        raise SystemExit("ยังไม่ได้ติดตั้ง openpyxl — รัน: pip install openpyxl")

    by_sid = {r["sid"]: r for r in roster}
    meta = {r["sid"]: r for r in results}

    wb = Workbook()
    wb.remove(wb.active)

    # ---- รายชื่อ ----
    ws = wb.create_sheet("รายชื่อ")
    head = ["#", "รหัสนักเรียน", "ชื่อ-สกุล", "ห้อง", "เลขที่", "สถานะ", "หมายเหตุ"]
    ws.append(head)
    _style_header(ws, 1, len(head))
    for i, r in enumerate(roster, 1):
        res = meta.get(r["sid"]) or {}
        status, note = _status(res)
        stu = res.get("student") or {}
        ws.append([i, r["sid"], stu.get("name") or r["name"],
                   stu.get("room") or r["room"], stu.get("seat") or r["seat"],
                   status, note])
    _set_widths(ws, [5, 12, 26, 8, 8, 14, 44])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    # ---- สรุป ----
    ws = wb.create_sheet("สรุป")
    head = ["#", "รหัส", "ชื่อ-สกุล", "ระดับชั้น", "ห้อง", "เลขที่", "ปี/ภาค",
            "GPA", "หน่วยกิตพื้นฐาน", "หน่วยกิตเพิ่มเติม", "หน่วยกิตรวม",
            "วิชาทั้งหมด", "เกรดออกแล้ว", "ปรับปรุง", "หมายเหตุ"]
    ws.append(head)
    _style_header(ws, 1, len(head))
    for i, res in enumerate(results, 1):
        stu = res.get("student") or {}
        rec = by_sid.get(res["sid"], {})
        for (y, t) in terms:
            d = res["terms"].get(f"{y}/{t}") or {}
            subs = d.get("subjects") or []
            cr = d.get("credits") or {}
            graded = sum(1 for s in subs
                         if (s.get("ระดับคะแนน") or "").strip() not in ("", "-", "0"))
            note = d.get("error") or ("ไม่มีข้อมูล" if not subs else "")
            row_i = ws.max_row + 1
            ws.append([i, res["sid"], stu.get("name") or rec.get("name", ""),
                       _label(d.get("level")), stu.get("room") or rec.get("room", ""),
                       stu.get("seat") or rec.get("seat", ""), f"ปี{y} ภาค{t}",
                       d.get("gpa"), cr.get("basic"), cr.get("extra"), cr.get("total"),
                       len(subs), graded, d.get("updated"), note])
            for col in (8, 9, 10, 11):
                ws.cell(row=row_i, column=col).number_format = "0.00"
    _set_widths(ws, [5, 12, 26, 10, 8, 8, 12, 8, 14, 14, 12, 10, 10, 20, 34])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    # ---- รวม (1 แถว = 1 วิชา) ----
    union = []
    for res in results:
        for d in res["terms"].values():
            for h in d.get("headers") or []:
                if h not in union:
                    union.append(h)
    ws = wb.create_sheet("รวม")
    head = ["รหัสนักเรียน", "ชื่อ-สกุล", "ห้อง", "เลขที่", "ปี/ภาค", *union, "ปรับปรุง"]
    ws.append(head)
    _style_header(ws, 1, len(head))
    for res in results:
        stu = res.get("student") or {}
        rec = by_sid.get(res["sid"], {})
        for (y, t) in terms:
            d = res["terms"].get(f"{y}/{t}") or {}
            for sub in d.get("subjects") or []:
                ws.append([res["sid"], stu.get("name") or rec.get("name", ""),
                           stu.get("room") or rec.get("room", ""),
                           stu.get("seat") or rec.get("seat", ""),
                           f"ปี{y} ภาค{t}",
                           *[sub.get(h) for h in union], d.get("updated")])
    widths = [12, 26, 8, 8, 12] + [14] * len(union) + [20]
    for j, h in enumerate(union):
        if h in ("รายวิชา", "ครูผู้สอน"):
            widths[5 + j] = 28
    _set_widths(ws, widths)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    # ---- ภาคละ 1 ชีต (1 แถว = 1 คน) ----
    for (y, t) in terms:
        key = f"{y}/{t}"
        codes, names = [], {}
        for res in results:
            for sub in (res["terms"].get(key) or {}).get("subjects") or []:
                code = sub.get("รหัสวิชา") or sub.get("รายวิชา")
                if not code:
                    continue
                if code not in codes:
                    codes.append(code)
                    names[code] = sub.get("รายวิชา") or ""
        if not codes:
            continue  # ภาคที่ยังไม่มีข้อมูล — ไม่ต้องสร้างชีตว่าง
        ws = wb.create_sheet(_safe_sheet(f"ปี{y}ภาค{t}"))
        ws.append(["#", "รหัส", "ชื่อ-สกุล", *codes, "GPA"])
        ws.append(["", "", "", *["" for _ in codes], ""])
        _style_header(ws, 1, 3 + len(codes) + 1)
        _style_header(ws, 2, 3 + len(codes) + 1)
        for j, code in enumerate(codes, start=4):
            ws.cell(row=2, column=j, value=names.get(code, ""))

        def seat_key(res):
            raw = (res.get("student") or {}).get("seat") or by_sid.get(res["sid"], {}).get("seat")
            try:
                return int(re.sub(r"\D", "", str(raw)) or 0)
            except Exception:
                return 0

        for i, res in enumerate(sorted(results, key=seat_key), 1):
            stu = res.get("student") or {}
            rec = by_sid.get(res["sid"], {})
            d = res["terms"].get(key) or {}
            grades = {}
            for sub in d.get("subjects") or []:
                code = sub.get("รหัสวิชา") or sub.get("รายวิชา")
                if code:
                    grades[code] = sub.get("ระดับคะแนน") or ""
            row_i = ws.max_row + 1
            ws.append([i, res["sid"], stu.get("name") or rec.get("name", ""),
                       *[grades.get(c, "") for c in codes], d.get("gpa")])
            ws.cell(row=row_i, column=4 + len(codes)).number_format = "0.00"

        _set_widths(ws, [5, 12, 26] + [11] * len(codes) + [8])
        ws.freeze_panes = "D3"
        for row in ws.iter_rows(min_row=3):
            for cell in row:
                cell.alignment = Alignment(vertical="top")

    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path


# ---------------- main ----------------
def main(argv=None):
    ap = argparse.ArgumentParser(description="ส่งออกผลการเรียนเป็น Excel")
    ap.add_argument("--roster", default=str(DEFAULT_ROSTER),
                    help="ไฟล์รายชื่อ CSV/XLSX (ค่าเริ่มต้น students.csv)")
    ap.add_argument("--term", default=None,
                    help="เลขปุ่ม 1-6 (0 = ทั้ง 6 ภาค); ไม่ใส่ = ถาม")
    ap.add_argument("--terms", default=None,
                    help="ปี-ภาค คั่นด้วย comma เช่น 1-1,2-1")
    ap.add_argument("--out", default="", help="ชื่อไฟล์ .xlsx ที่ต้องการ")
    ap.add_argument("--limit", type=int, default=0, help="ทำแค่ N คนแรก")
    ap.add_argument("--workers", type=int, default=3, help="จำนวนคนที่ล็อกอินพร้อมกัน")
    ap.add_argument("--stagger", type=float, default=2, help="หน่วงระหว่างชุด (วินาที)")
    ap.add_argument("--retries", type=int, default=2,
                    help="ลองใหม่กี่รอบเมื่อเจอเว็บ 5xx/หมดเวลา (ค่าเริ่มต้น 2)")
    ap.add_argument("--headed", action="store_true", help="เปิดหน้าจอเบราว์เซอร์ให้ดู")
    ap.add_argument("--dry-run", action="store_true", help="ตรวจรายชื่อแล้วจบ ไม่เปิดเว็บ")
    ap.add_argument("--save-html", action="store_true", help="เก็บ HTML ไว้ debug parser")
    ap.add_argument("--shots", action="store_true", help="เก็บภาพหน้าจอแต่ละภาค")
    ap.add_argument("--verbose", action="store_true", help="แสดงทุกขั้นตอนของเอนจิน")
    args = ap.parse_args(argv)

    if args.term is not None and args.terms:
        raise SystemExit("เลือกอย่างใดอย่างหนึ่ง: --term หรือ --terms")
    if args.term is not None:
        terms = parse_term_arg(args.term)
    elif args.terms:
        terms = parse_terms(args.terms)
    else:
        terms = choose_term()

    roster, errors = load_roster(args.roster)
    total = len(roster)

    if args.limit > 0:
        roster = roster[:args.limit]

    if args.dry_run:
        print_roster(roster, errors, terms, total)
        return 1 if errors else 0

    if errors:
        print(f"[!] ข้าม {len(errors)} แถวที่ใช้ไม่ได้:")
        for e in errors:
            print("   -", e)
    if not roster:
        raise SystemExit("ไม่มีรายชื่อที่ใช้ได้ — แก้ไฟล์แล้วลองใหม่")

    if not args.verbose:
        g.step = lambda *a, **k: None   # เงียบ: ไม่พิมพ์ทุกขั้นของ login

    print(f"เริ่มดึง {len(roster)} คน | " + ", ".join(f"ปี{y}ภาค{t}" for y, t in terms)
          + f" | workers={args.workers}")
    t0 = time.time()
    try:
        results = asyncio.run(_run(roster, terms, args))
    except KeyboardInterrupt:
        print("\nยกเลิกกลางทาง — ยังไม่ได้เขียนไฟล์")
        return 130

    out = Path(args.out) if args.out else (
        EXPORT_DIR / f"grades_{datetime.datetime.now():%Y%m%d_%H%M}.xlsx")
    out = write_workbook(results, roster, terms, out)

    ok = sum(1 for r in results if _status(r)[0] == "ok")
    print(f"\nเสร็จใน {time.time() - t0:.0f}s | สำเร็จ {ok}/{len(results)}")
    print(f"บันทึก: {out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
